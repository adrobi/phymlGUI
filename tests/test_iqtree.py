"""Tests for the IQ-TREE engine, ModelFinder and the XLSX summary."""
import os, sys, time, shutil, re, json

os.environ['QT_QPA_PLATFORM'] = 'offscreen'
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.chdir(ROOT)

import phymlgui
from PySide6.QtWidgets import QMessageBox

IQTREE_EXE = os.path.join(ROOT, 'tools', 'iqtree2', 'iqtree2.exe')
if not os.path.exists(IQTREE_EXE):
    print('SKIP: IQ-TREE binary not available')
    sys.exit(0)

app = phymlgui.QApplication([])

# 1. worker-level: MFP + ultrafast bootstrap on pre-aligned synthetic data
run_dir = os.path.join('testrun', 't10_iqtree')
shutil.rmtree(run_dir, ignore_errors=True)
os.makedirs(run_dir, exist_ok=True)
shutil.copy('testrun/test_prealigned.fasta', run_dir)
os.chdir(run_dir)

config = {
    'engine': 'iqtree', 'iqtree_path': IQTREE_EXE,
    'phyml_path': os.path.join(ROOT, 'phymlgui v1.1', 'phyml-mpi.exe'),
    'muscle_path': os.path.join(ROOT, 'tools', 'muscle.exe'),
    'use_mpi': False, 'mpi_np': '2', 'datatype': 'nt', 'model': 'MFP',
    'support_method': 'bootstrap', 'bootstrap_replicates': '1000',
    'freq': 'default', 'tstv': 'e', 'pinv': '0.0',
    'classes': '4', 'gamma': 'e', 'search': 'NNI', 'optimize': 'tlr', 'parsimony': False,
}
logs, done = [], []
worker = phymlgui.PipelineWorker(config, os.path.abspath('test_prealigned.fasta'))
worker.log_signal.connect(logs.append)
worker.finished_signal.connect(lambda ok, tree: done.append((ok, tree)))
worker.run()
os.chdir(ROOT)

assert done[0][0], 'iqtree pipeline failed:\n' + '\n'.join(logs[-8:])
tree_rel = os.path.relpath(done[0][1], ROOT)
assert os.path.exists(done[0][1]), done[0][1]
from Bio import Phylo
tree = Phylo.read(done[0][1], 'newick')
assert len(tree.get_terminals()) == 8
row = worker.results[0]
assert row['model'] and row['model'] != 'n/a', row
iq_report = open(os.path.join(run_dir, 'test_prealigned_iqtree.iqtree'), encoding='utf-8').read()
lnl_file = re.search(r'Log-likelihood of the tree:\s*(-?[\d.]+)', iq_report).group(1)
assert row['lnL'] == lnl_file, (row['lnL'], lnl_file)
print(f'PASS iqtree engine: MFP picked {row["model"]}, lnL {row["lnL"]}, 8 taxa')

# 1b. PhyML-sized replicate count (100) must not break UFBoot (needs >= 1000)
os.chdir(run_dir)
logs, done = [], []
worker = phymlgui.PipelineWorker(dict(config, bootstrap_replicates='100'),
                                 os.path.abspath('test_prealigned.fasta'))
worker.log_signal.connect(logs.append)
worker.finished_signal.connect(lambda ok, tree: done.append((ok, tree)))
worker.run()
os.chdir(ROOT)
assert done[0][0], 'iqtree with 100 replicates failed:\n' + '\n'.join(logs[-8:])
assert any(' -B 1000' in l for l in logs), 'UFBoot replicates were not raised to 1000'
print('PASS iqtree ufboot: 100 replicates raised to 1000, run succeeds')

# 2. GUI: engine validation, model list, full run + xlsx summary
cfg_bak = phymlgui.CONFIG_FILE + '.bak'
had_cfg = os.path.exists(phymlgui.CONFIG_FILE)
if had_cfg:
    shutil.copy(phymlgui.CONFIG_FILE, cfg_bak)
try:
    w = phymlgui.PhyMLApp()
    warned = []
    phymlgui.QMessageBox.warning = staticmethod(lambda *a, **k: warned.append(a) or QMessageBox.Ok)
    w.engine_setting = 'iqtree'
    w.update_models()
    assert w.cb_model.currentData() == 'MFP'
    w.iqtree_path_input.setText('')  # no binary -> blocked with warning
    w.fasta_path_input.setText(os.path.abspath(os.path.join(run_dir, 'test_prealigned.fasta')))
    w.start_pipeline()
    assert warned and (w.worker is None or not w.worker.isRunning())

    w.iqtree_path_input.setText(IQTREE_EXE)
    w.set_combo_by_data(w.cb_support_method, 'bootstrap')
    w.le_bootstrap_replicates.setText('100')
    w.start_pipeline()
    guard = 0
    while not w.btn_run.isEnabled() and guard < 400:
        app.processEvents(); time.sleep(0.3); guard += 1
    assert w.btn_run.isEnabled(), 'GUI iqtree run did not finish'
    assert w.worker.results[0]['status'] == 'OK', \
        'GUI iqtree run failed:\n' + w.log_console.toPlainText()[-800:]
    xlsx = os.path.join(run_dir, 'phymlgui_summary.xlsx')
    assert os.path.exists(xlsx), 'xlsx summary missing'
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    assert wb.active.max_row == 2
    print('PASS iqtree GUI: validation, model list, run + xlsx summary')

    # 3. every model offered in the IQ-TREE model list is accepted by IQ-TREE
    import random
    import subprocess
    rnd = random.Random(7)
    aa_phy = os.path.join(run_dir, 'models_aa.phy')
    with open(aa_phy, 'w') as f:
        base = ''.join(rnd.choice('ACDEFGHIKLMNPQRSTVWY') for _ in range(120))
        f.write('6 120\n')
        for i in range(6):
            seq = ''.join(c if rnd.random() > 0.15 else rnd.choice('ACDEFGHIKLMNPQRSTVWY')
                          for c in base)
            f.write(f'P{i} {seq}\n')
    alignments = {'nt': os.path.join(run_dir, 'test_prealigned.phy'), 'aa': aa_phy}
    rejected = []
    for dt, phy in alignments.items():
        w.set_combo_by_data(w.cb_datatype, dt)
        models = [w.cb_model.itemData(i) for i in range(w.cb_model.count())]
        for m in models:
            if m == 'MFP':
                continue  # ModelFinder is exercised by the runs above
            r = subprocess.run([IQTREE_EXE, '-s', phy, '-m', m, '-n', '0', '-T', '1', '-redo',
                                '--prefix', os.path.join(run_dir, 'model_probe')],
                               capture_output=True, text=True)
            if r.returncode != 0:
                rejected.append(m)
    assert not rejected, f'IQ-TREE rejects listed models: {rejected}'
    print('PASS iqtree models: every listed nt/aa model is accepted by IQ-TREE')
finally:
    if had_cfg:
        shutil.move(cfg_bak, phymlgui.CONFIG_FILE)
print('ALL IQ-TREE TESTS PASSED')
