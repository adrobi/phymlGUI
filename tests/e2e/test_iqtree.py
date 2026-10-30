"""IQ-TREE engine, ModelFinder and the XLSX summary."""
import random
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from Bio import AlignIO, Phylo

from datasets import synthetic  # noqa: F401  (fixture)
from phyml_core.engines import ENGINES
from phyml_core.pipeline import PipelineWorker

pytestmark = pytest.mark.needs_iqtree


@pytest.fixture
def run_dir(tmp_path, synthetic):
    """A working folder holding the pre-aligned synthetic dataset."""
    folder = tmp_path / 't10_iqtree'
    folder.mkdir()
    shutil.copy(synthetic['aligned'], folder / 'test_prealigned.fasta')
    return folder


def iqtree_config(binaries, **overrides):
    config = {
        'engine': 'iqtree', 'iqtree_path': binaries['iqtree'],
        'use_mpi': False, 'mpi_np': '2', 'datatype': 'nt', 'model': 'MFP',
        'support_method': 'bootstrap', 'bootstrap_replicates': '1000',
        'freq': 'default', 'tstv': 'e', 'pinv': '0.0',
        'classes': '4', 'gamma': 'e', 'search': 'NNI', 'optimize': 'tlr', 'parsimony': False,
    }
    config.update(overrides)
    return config


def run_worker(config, fasta):
    """Run the worker synchronously; returns (worker, logs, finished results)."""
    logs, done = [], []
    worker = PipelineWorker(config, str(fasta))
    worker.log_signal.connect(logs.append)
    worker.finished_signal.connect(lambda ok, tree: done.append((ok, tree)))
    worker.run()
    return worker, logs, done


def test_mfp_ufboot_run(qapp, monkeypatch, binaries, run_dir):
    """Worker-level: ModelFinder + ultrafast bootstrap on pre-aligned synthetic data."""
    monkeypatch.chdir(run_dir)
    worker, logs, done = run_worker(iqtree_config(binaries), run_dir / 'test_prealigned.fasta')

    assert done[0][0], 'iqtree pipeline failed:\n' + '\n'.join(logs[-8:])
    tree_path = done[0][1]
    assert Path(tree_path).exists(), tree_path
    tree = Phylo.read(tree_path, 'newick')
    assert len(tree.get_terminals()) == 8
    row = worker.results[0]
    assert row['model'] and row['model'] != 'n/a', row
    iq_report = (run_dir / 'test_prealigned_iqtree.iqtree').read_text(encoding='utf-8')
    lnl_file = re.search(r'Log-likelihood of the tree:\s*(-?[\d.]+)', iq_report).group(1)
    assert row['lnL'] == lnl_file, (row['lnL'], lnl_file)


def test_100_replicates_raised_to_1000(qapp, monkeypatch, binaries, run_dir):
    """A PhyML-sized replicate count (100) must not break UFBoot (needs >= 1000)."""
    monkeypatch.chdir(run_dir)
    config = iqtree_config(binaries, bootstrap_replicates='100')
    _, logs, done = run_worker(config, run_dir / 'test_prealigned.fasta')
    assert done[0][0], 'iqtree with 100 replicates failed:\n' + '\n'.join(logs[-8:])
    assert any(' -B 1000' in line for line in logs), 'UFBoot replicates were not raised to 1000'


def test_gui_validation_and_run_with_xlsx(window, warned, qtbot, binaries, run_dir):
    """GUI: engine validation, model list, full run + xlsx summary."""
    window.engine_setting = 'iqtree'
    window.update_models()
    assert window.cb_model.currentData() == 'MFP'
    window.iqtree_path_input.setText('')  # no binary -> blocked with warning
    window.fasta_path_input.setText(str(run_dir / 'test_prealigned.fasta'))
    window.start_pipeline()
    assert warned and (window.worker is None or not window.worker.isRunning())

    window.iqtree_path_input.setText(binaries['iqtree'])
    window.set_combo_by_data(window.cb_support_method, 'bootstrap')
    window.le_bootstrap_replicates.setText('100')
    window.start_pipeline()
    qtbot.waitUntil(lambda: window.btn_run.isEnabled(), timeout=180_000)
    window.worker.wait()
    assert window.worker.results[0]['status'] == 'OK', \
        'GUI iqtree run failed:\n' + window.log_console.toPlainText()[-800:]
    xlsx = run_dir / 'phymlgui_summary.xlsx'
    assert xlsx.exists(), 'xlsx summary missing'
    from openpyxl import load_workbook  # optional dependency of the app
    wb = load_workbook(str(xlsx))
    assert wb.active.max_row == 2


def test_every_listed_model_accepted(binaries, run_dir):
    """Every model the IQ-TREE engine offers is accepted by IQ-TREE."""
    nt_phy = run_dir / 'test_prealigned.phy'
    AlignIO.convert(str(run_dir / 'test_prealigned.fasta'), 'fasta',
                    str(nt_phy), 'phylip-relaxed')
    rnd = random.Random(7)
    aa_phy = run_dir / 'models_aa.phy'
    with open(aa_phy, 'w') as f:
        base = ''.join(rnd.choice('ACDEFGHIKLMNPQRSTVWY') for _ in range(120))
        f.write('6 120\n')
        for i in range(6):
            seq = ''.join(c if rnd.random() > 0.15 else rnd.choice('ACDEFGHIKLMNPQRSTVWY')
                          for c in base)
            f.write(f'P{i} {seq}\n')

    rejected = []
    for dt, phy in {'nt': nt_phy, 'aa': aa_phy}.items():
        for m in ENGINES['iqtree'].models(dt):
            if m == 'MFP':
                continue  # ModelFinder is exercised by the runs above
            r = subprocess.run([binaries['iqtree'], '-s', str(phy), '-m', m, '-n', '0',
                                '-T', '1', '-redo', '--prefix', str(run_dir / 'model_probe')],
                               capture_output=True, text=True, timeout=120)
            if r.returncode != 0:
                rejected.append(m)
    assert not rejected, f'IQ-TREE rejects listed models: {rejected}'
